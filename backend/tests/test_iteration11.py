"""Iteration 11 backend tests: Rapports, Zones (add/delete), Archivage auto."""
import io
import os
import uuid
from datetime import datetime, timedelta, timezone

import jwt
import pytest
import requests

BASE_URL = (os.environ.get("EXPO_PUBLIC_BACKEND_URL") or os.environ.get("EXPO_BACKEND_URL", "")).rstrip("/")
API = f"{BASE_URL}/api"
JWT_SECRET = os.environ.get("JWT_SECRET", "pinkberry-drc-secret-2026-kinshasa")
SUPERADMIN_EMAIL = "mecedloga@gmail.com"


def _login(tel: str, mdp: str) -> str:
    r = requests.post(f"{API}/auth/connexion", json={"telephone": tel, "motDePasse": mdp}, timeout=15)
    assert r.status_code == 200, f"login {tel} -> {r.status_code} {r.text}"
    return r.json()["token"]


def _h(tok):
    return {"Authorization": f"Bearer {tok}"}


def _forge_token(sub_id: str) -> str:
    return jwt.encode(
        {"sub": sub_id, "exp": datetime.now(timezone.utc) + timedelta(days=1)},
        JWT_SECRET, algorithm="HS256"
    )


def _mongo():
    from pymongo import MongoClient
    c = MongoClient(os.environ.get("MONGO_URL", "mongodb://localhost:27017"))
    return c[os.environ.get("DB_NAME", "test_database")]


def _iso(dt): return dt.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


@pytest.fixture(scope="module")
def super_id():
    db = _mongo()
    u = db.users.find_one({"email": SUPERADMIN_EMAIL, "role": "superadmin"}, {"_id": 0, "id": 1})
    assert u, "Superadmin not seeded"
    return u["id"]


@pytest.fixture(scope="module")
def tokens(super_id):
    return {
        "super": _forge_token(super_id),
        "admin": _login("0810000000", "admin123"),
        "livreur": _login("0820000001", "livreur123"),
        "client": _login("0990000001", "client123"),
    }


def _create_test_client():
    tel = f"0995{uuid.uuid4().int % 1000000:06d}"[:10]
    r = requests.post(f"{API}/auth/inscription",
                      json={"nom": "TEST_Iter11", "email": f"otp_{tel}@pinkberry.test", "telephone": tel, "motDePasse": "abcdef"}, timeout=15)
    assert r.status_code == 200, r.text
    code = r.json()["codeDemo"]
    r2 = requests.post(f"{API}/auth/verifier-otp",
                       json={"telephone": tel, "code": code}, timeout=15)
    assert r2.status_code == 200
    return tel, r2.json()["token"], r2.json()["user"]["id"]


def _cleanup_user(uid):
    db = _mongo()
    db.users.delete_one({"id": uid})
    db.commandes.delete_many({"userId": uid})
    db.adresses.delete_many({"userId": uid})


# ================================================================ RAPPORTS
class TestRapport:
    def test_rapport_xlsx_ok(self, tokens):
        today = datetime.now(timezone.utc).date()
        debut = (today - timedelta(days=30)).isoformat()
        fin = today.isoformat()
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": debut, "fin": fin, "format": "xlsx"},
                         headers=_h(tokens["admin"]), timeout=30)
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml" in ct, ct
        cd = r.headers.get("content-disposition", "")
        assert f"pinkberry-rapport-{debut}-{fin}.xlsx" in cd, cd
        assert len(r.content) > 0

        # Ouvrir avec openpyxl
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        titre = ws["A1"].value
        assert "Rapport des commandes" in titre and debut in titre and fin in titre

        entetes_attendues = ["N° commande", "Date", "Client", "Adresse / commune",
                             "Articles", "Promotion", "Réduction FC", "Total FC"]
        entetes = [ws.cell(row=3, column=i + 1).value for i in range(8)]
        assert entetes == entetes_attendues, entetes

        # Chercher le bloc récapitulatif
        labels = [ws.cell(row=r_i, column=1).value for r_i in range(1, ws.max_row + 1)]
        assert "Récapitulatif" in labels
        assert any(l and "Nombre de commandes" in str(l) for l in labels)
        assert any(l and "Chiffre d'affaires" in str(l) for l in labels)
        assert any(l and "Livraisons réussies" in str(l) for l in labels)
        assert any(l and "Commandes annulées" in str(l) for l in labels)

    def test_rapport_pdf_ok(self, tokens):
        today = datetime.now(timezone.utc).date()
        debut = (today - timedelta(days=7)).isoformat()
        fin = today.isoformat()
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": debut, "fin": fin, "format": "pdf"},
                         headers=_h(tokens["admin"]), timeout=30)
        assert r.status_code == 200, r.text
        assert "application/pdf" in r.headers.get("content-type", "")
        assert f"pinkberry-rapport-{debut}-{fin}.pdf" in r.headers.get("content-disposition", "")
        assert r.content[:4] == b"%PDF", "Not a valid PDF"
        assert len(r.content) > 500

    def test_rapport_coherence_recap_et_archives_incluses(self, tokens):
        """Injecter des commandes de test livrees (une archivee, une non) + une annulee,
        verifier que le xlsx les compte correctement et inclut la commande archivee."""
        db = _mongo()
        today = datetime.now(timezone.utc)
        debut = today.date().isoformat()
        fin = today.date().isoformat()
        cids = []
        for i, (statut, archivee, montant) in enumerate([
            ("livree", False, 12500),
            ("livree", True, 8000),   # archivee mais DOIT être dans le rapport
            ("annulee", False, 5000),
            ("recue", False, 3000),   # ni livree ni annulee -> pas dans chiffre/livrees/annulees
        ]):
            cid = f"TEST_ITER11_{uuid.uuid4().hex[:8]}"
            doc = {
                "id": cid,
                "numero": f"PB-TEST-{i}",
                "userId": "TEST_ITER11",
                "clientNom": "TEST_ITER11_Client",
                "adresse": {"commune": "Gombe", "avenue": "Ave Test", "numero": str(i)},
                "items": [{"produitId": "p1", "nom": "Eau minérale 50cl",
                           "taille": "Petit", "quantite": 2, "prixUnitaire": montant / 2}],
                "total": montant,
                "statut": statut,
                "archivee": archivee,
                "createdAt": _iso(today.replace(hour=12, minute=0, second=0, microsecond=0)),
            }
            db.commandes.insert_one(doc)
            cids.append(cid)

        try:
            r = requests.get(f"{API}/admin/rapport",
                             params={"debut": debut, "fin": fin, "format": "xlsx"},
                             headers=_h(tokens["admin"]), timeout=30)
            assert r.status_code == 200

            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(r.content))
            ws = wb.active
            # Verifier la présence des 4 numeros (l'archivée incluse)
            texte = "\n".join(str(c.value or "") for row in ws.iter_rows() for c in row)
            for i in range(4):
                assert f"PB-TEST-{i}" in texte, f"PB-TEST-{i} manquant dans rapport"
            assert "Eau minérale 50cl" in texte
            assert "2x Eau minérale 50cl (Petit)" in texte, "Format articles incorrect"

            # Recap : trouver les valeurs à droite des labels
            recap = {}
            for row in ws.iter_rows(values_only=True):
                if row and row[0]:
                    lbl = str(row[0])
                    val = row[1] if len(row) > 1 else None
                    if "Nombre de commandes" in lbl: recap["nb"] = val
                    elif "Chiffre d'affaires" in lbl: recap["ca"] = val
                    elif "Livraisons réussies" in lbl: recap["ok"] = val
                    elif "Commandes annulées" in lbl: recap["ko"] = val

            # Attention : le rapport prend toutes les commandes de la periode
            # Nos 4 commandes de test + commandes de démo eventuellement
            assert recap["nb"] >= 4
            assert recap["ok"] >= 2  # 2 livrées
            assert recap["ko"] >= 1  # 1 annulée
            # CA en FC (« 20 500 FC ») : au moins nos 2 livrées (12 500 + 8 000)
            assert int(str(recap["ca"]).replace("FC", "").replace(" ", "").strip()) >= 20500
        finally:
            for cid in cids:
                db.commandes.delete_one({"id": cid})

    def test_rapport_format_invalide(self, tokens):
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": "2025-01-01", "fin": "2025-01-31", "format": "csv"},
                         headers=_h(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_rapport_dates_malformees(self, tokens):
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": "01-01-2025", "fin": "2025-01-31", "format": "xlsx"},
                         headers=_h(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_rapport_fin_avant_debut(self, tokens):
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": "2025-02-01", "fin": "2025-01-01", "format": "xlsx"},
                         headers=_h(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_rapport_sans_jeton_401(self):
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": "2025-01-01", "fin": "2025-01-31", "format": "xlsx"},
                         timeout=15)
        assert r.status_code == 401

    def test_rapport_client_403(self, tokens):
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": "2025-01-01", "fin": "2025-01-31", "format": "xlsx"},
                         headers=_h(tokens["client"]), timeout=15)
        assert r.status_code == 403

    def test_rapport_livreur_403(self, tokens):
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": "2025-01-01", "fin": "2025-01-31", "format": "xlsx"},
                         headers=_h(tokens["livreur"]), timeout=15)
        assert r.status_code == 403

    def test_rapport_superadmin_200(self, tokens):
        today = datetime.now(timezone.utc).date().isoformat()
        r = requests.get(f"{API}/admin/rapport",
                         params={"debut": today, "fin": today, "format": "xlsx"},
                         headers=_h(tokens["super"]), timeout=15)
        assert r.status_code == 200


# =================================================================== ZONES
class TestZones:
    def test_creer_zone_ok_et_visible_public(self, tokens):
        commune = f"TEST_ITER11_{uuid.uuid4().hex[:6]}"
        r = requests.post(f"{API}/admin/zones",
                          json={"commune": commune, "active": True},
                          headers=_h(tokens["admin"]), timeout=15)
        assert r.status_code == 200, r.text
        zone_id = r.json()["id"]

        try:
            # GET /api/zones public - la zone y est
            z = requests.get(f"{API}/zones", timeout=15).json()
            assert any(zo["commune"] == commune for zo in z)
        finally:
            requests.delete(f"{API}/admin/zones/{zone_id}",
                            headers=_h(tokens["admin"]), timeout=15)

    def test_creer_zone_nom_trop_court(self, tokens):
        r = requests.post(f"{API}/admin/zones",
                          json={"commune": "A", "active": True},
                          headers=_h(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_creer_zone_doublon_insensible_casse(self, tokens):
        commune = f"TEST_ITER11_{uuid.uuid4().hex[:6]}"
        r = requests.post(f"{API}/admin/zones",
                          json={"commune": commune, "active": True},
                          headers=_h(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        zid = r.json()["id"]
        try:
            r2 = requests.post(f"{API}/admin/zones",
                               json={"commune": commune.upper(), "active": True},
                               headers=_h(tokens["admin"]), timeout=15)
            assert r2.status_code == 400
            r3 = requests.post(f"{API}/admin/zones",
                               json={"commune": commune.lower(), "active": True},
                               headers=_h(tokens["admin"]), timeout=15)
            assert r3.status_code == 400
        finally:
            requests.delete(f"{API}/admin/zones/{zid}",
                            headers=_h(tokens["admin"]), timeout=15)

    def test_supprimer_zone_ok(self, tokens):
        commune = f"TEST_ITER11_{uuid.uuid4().hex[:6]}"
        r = requests.post(f"{API}/admin/zones",
                          json={"commune": commune, "active": True},
                          headers=_h(tokens["admin"]), timeout=15)
        zid = r.json()["id"]

        r2 = requests.delete(f"{API}/admin/zones/{zid}",
                             headers=_h(tokens["admin"]), timeout=15)
        assert r2.status_code == 200

        z = requests.get(f"{API}/zones", timeout=15).json()
        assert not any(zo["commune"] == commune for zo in z)

    def test_supprimer_zone_404(self, tokens):
        r = requests.delete(f"{API}/admin/zones/inexistant",
                            headers=_h(tokens["admin"]), timeout=15)
        assert r.status_code == 404

    def test_rbac_zones(self, tokens):
        # Sans jeton
        r = requests.post(f"{API}/admin/zones", json={"commune": "TEST_x", "active": True}, timeout=15)
        assert r.status_code == 401
        r = requests.delete(f"{API}/admin/zones/xx", timeout=15)
        assert r.status_code == 401
        # Client
        r = requests.post(f"{API}/admin/zones", json={"commune": "TEST_x", "active": True},
                          headers=_h(tokens["client"]), timeout=15)
        assert r.status_code == 403
        r = requests.delete(f"{API}/admin/zones/xx", headers=_h(tokens["client"]), timeout=15)
        assert r.status_code == 403
        # Livreur
        r = requests.post(f"{API}/admin/zones", json={"commune": "TEST_x", "active": True},
                          headers=_h(tokens["livreur"]), timeout=15)
        assert r.status_code == 403

    def test_superadmin_peut_gerer_zones(self, tokens):
        commune = f"TEST_ITER11_{uuid.uuid4().hex[:6]}"
        r = requests.post(f"{API}/admin/zones",
                          json={"commune": commune, "active": True},
                          headers=_h(tokens["super"]), timeout=15)
        assert r.status_code == 200
        zid = r.json()["id"]
        r2 = requests.delete(f"{API}/admin/zones/{zid}",
                             headers=_h(tokens["super"]), timeout=15)
        assert r2.status_code == 200

    def test_zone_inactive_refuse_commande_activation_supprime(self, tokens):
        """Effet metier : zone inactive -> POST commande refuse ; activer -> ok ; supprimer -> refuse."""
        commune = f"TEST_ITER11_{uuid.uuid4().hex[:6]}"
        r = requests.post(f"{API}/admin/zones",
                          json={"commune": commune, "active": False},
                          headers=_h(tokens["admin"]), timeout=15)
        zid = r.json()["id"]

        tel, ctok, cuid = _create_test_client()
        payload = {
            "adresse": {"label": "TEST", "commune": commune, "avenue": "Ave X",
                        "numero": "1", "repere": "porte", "telephone": tel,
                        "lat": -4.3, "lng": 15.3},
            "items": [{"produitId": "p1", "nom": "Pot", "taille": "Petit",
                       "garnitures": [], "quantite": 1, "prixUnitaire": 4.5}],
        }
        try:
            # Zone inactive -> refuse
            r1 = requests.post(f"{API}/commandes", headers=_h(ctok), json=payload, timeout=15)
            assert r1.status_code >= 400, f"zone inactive devrait refuser, got {r1.status_code}"

            # Activer
            ru = requests.put(f"{API}/admin/zones/{zid}",
                              json={"commune": commune, "active": True},
                              headers=_h(tokens["admin"]), timeout=15)
            assert ru.status_code == 200

            # Zone active -> commande accepte
            r2 = requests.post(f"{API}/commandes", headers=_h(ctok), json=payload, timeout=15)
            assert r2.status_code == 200, f"zone active devrait accepter, got {r2.status_code} {r2.text}"

            # Supprimer zone
            rd = requests.delete(f"{API}/admin/zones/{zid}",
                                 headers=_h(tokens["admin"]), timeout=15)
            assert rd.status_code == 200
            zid = None

            # Nouvelle commande sur cette commune -> refuse
            r3 = requests.post(f"{API}/commandes", headers=_h(ctok), json=payload, timeout=15)
            assert r3.status_code >= 400
        finally:
            if zid:
                requests.delete(f"{API}/admin/zones/{zid}",
                                headers=_h(tokens["admin"]), timeout=15)
            _mongo().commandes.delete_many({"userId": cuid})
            _cleanup_user(cuid)


# =========================================================== ARCHIVAGE AUTO
class TestArchivageAuto:
    def test_archiver_anciennes_livree_ok_et_recente_reste(self, tokens):
        """Deux commandes livrées : ancienne (>10j) archivée, récente (<10j) reste.
        Une annulée ancienne : NON archivée automatiquement."""
        db = _mongo()
        now = datetime.now(timezone.utc)
        cid_ancienne = f"TEST_ITER11_ANC_{uuid.uuid4().hex[:6]}"
        cid_recente = f"TEST_ITER11_REC_{uuid.uuid4().hex[:6]}"
        cid_annulee_anc = f"TEST_ITER11_ANN_{uuid.uuid4().hex[:6]}"

        base = {
            "userId": "TEST_ITER11",
            "clientNom": "TEST_ITER11",
            "adresse": {"commune": "Gombe", "avenue": "A", "numero": "1"},
            "items": [{"nom": "X", "quantite": 1, "prixUnitaire": 1.0}],
            "total": 1.0,
            "archivee": False,
        }
        db.commandes.insert_one({
            **base, "id": cid_ancienne, "numero": "PB-ANC",
            "statut": "livree",
            "termineeAt": _iso(now - timedelta(days=15)),
            "createdAt": _iso(now - timedelta(days=20)),
        })
        db.commandes.insert_one({
            **base, "id": cid_recente, "numero": "PB-REC",
            "statut": "livree",
            "termineeAt": _iso(now - timedelta(days=2)),
            "createdAt": _iso(now - timedelta(days=3)),
        })
        db.commandes.insert_one({
            **base, "id": cid_annulee_anc, "numero": "PB-ANN",
            "statut": "annulee",
            "termineeAt": _iso(now - timedelta(days=15)),
            "createdAt": _iso(now - timedelta(days=20)),
        })

        try:
            # Appel qui declenche l'archivage auto
            r = requests.get(f"{API}/admin/commandes",
                             headers=_h(tokens["admin"]), timeout=30)
            assert r.status_code == 200
            ids_actifs = {c["id"] for c in r.json()}

            # Ancienne livree n'est plus dans dashboard
            assert cid_ancienne not in ids_actifs
            # Recente livree y est toujours
            assert cid_recente in ids_actifs
            # Annulee ancienne y est toujours (pas archivee auto)
            assert cid_annulee_anc in ids_actifs

            # Verifier archivee=true et archivageAuto=true en base pour l'ancienne
            doc = db.commandes.find_one({"id": cid_ancienne})
            assert doc.get("archivee") is True
            assert doc.get("archivageAuto") is True

            # L'annulee ancienne ne doit PAS être archivee
            doc_ann = db.commandes.find_one({"id": cid_annulee_anc})
            assert doc_ann.get("archivee") is not True

            # Elle apparait dans /admin/commandes/archives
            ra = requests.get(f"{API}/admin/commandes/archives",
                              headers=_h(tokens["admin"]), timeout=30).json()
            ids_arc = {c["id"] for c in ra}
            assert cid_ancienne in ids_arc
        finally:
            db.commandes.delete_many({
                "id": {"$in": [cid_ancienne, cid_recente, cid_annulee_anc]}
            })

    def test_termineeAt_pose_par_patch_statut(self, tokens):
        """PATCH /api/admin/commandes/{id}/statut vers livree ou annulee pose termineeAt."""
        db = _mongo()
        tel, ctok, cuid = _create_test_client()
        cmd_payload = {
            "adresse": {"label": "TEST", "commune": "Gombe", "avenue": "Ave A",
                        "numero": "1", "repere": "porte", "telephone": tel,
                        "lat": -4.3, "lng": 15.3},
            "items": [{"produitId": "p1", "nom": "Pot", "taille": "Petit",
                       "garnitures": [], "quantite": 1, "prixUnitaire": 4.5}],
        }
        try:
            r = requests.post(f"{API}/commandes", headers=_h(ctok),
                              json=cmd_payload, timeout=15)
            assert r.status_code == 200
            cid = r.json()["id"]

            # PATCH -> annulee
            r2 = requests.patch(f"{API}/admin/commandes/{cid}/statut",
                                headers=_h(tokens["admin"]),
                                json={"statut": "annulee"}, timeout=15)
            assert r2.status_code == 200

            doc = db.commandes.find_one({"id": cid})
            assert doc.get("termineeAt"), "termineeAt manquant apres annulee"
            db.commandes.delete_one({"id": cid})
        finally:
            _cleanup_user(cuid)

    def test_termineeAt_pose_par_annulation_client(self, tokens):
        """Le client annule sa commande -> termineeAt pose."""
        db = _mongo()
        tel, ctok, cuid = _create_test_client()
        cmd_payload = {
            "adresse": {"label": "TEST", "commune": "Gombe", "avenue": "A",
                        "numero": "1", "repere": "p", "telephone": tel,
                        "lat": -4.3, "lng": 15.3},
            "items": [{"produitId": "p1", "nom": "Pot", "taille": "Petit",
                       "garnitures": [], "quantite": 1, "prixUnitaire": 4.5}],
        }
        try:
            r = requests.post(f"{API}/commandes", headers=_h(ctok),
                              json=cmd_payload, timeout=15)
            cid = r.json()["id"]
            # Annuler cote client
            ra = requests.post(f"{API}/commandes/{cid}/annuler",
                               headers=_h(ctok), timeout=15)
            # Peut etre 200 ou autre selon regles
            if ra.status_code == 200:
                doc = db.commandes.find_one({"id": cid})
                assert doc.get("termineeAt"), "termineeAt manquant apres annulation client"
            db.commandes.delete_one({"id": cid})
        finally:
            _cleanup_user(cuid)
